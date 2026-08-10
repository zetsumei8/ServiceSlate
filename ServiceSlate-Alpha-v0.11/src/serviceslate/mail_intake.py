from __future__ import annotations

import csv
import hashlib
import html
import io
import json
import mimetypes
import re
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from email import policy
from email.parser import BytesParser
from email.utils import parseaddr
from pathlib import Path
from typing import Any, Iterable

from .db import FILES_DIR, connect, new_id, utcnow


@dataclass(slots=True)
class MailAttachment:
    name: str
    content_type: str | None
    data: bytes


@dataclass(slots=True)
class MailEnvelope:
    external_id: str
    internet_message_id: str | None
    sender: str | None
    subject: str
    received_at: str | None
    body: str
    attachments: list[MailAttachment]


# These aliases are intentionally broad only for staging. ServiceSlate never commits a
# staged value without matching/review rules in the normal import workflow.
ALIASES: dict[str, tuple[str, ...]] = {
    "customer": (
        "customer", "customer name", "account", "account name", "company", "company name",
        "business", "business name", "dealer", "dealership", "client", "client name",
    ),
    "location": (
        "location", "site", "site name", "service location", "store", "branch", "facility",
    ),
    "source_reference": (
        "work order", "work order number", "work order #", "work_order", "wo", "wo #", "job",
        "job number", "job #", "reference", "reference number", "submission id", "submission #",
        "form submission id", "record id",
    ),
    "performed_at": (
        "date", "work date", "completed date", "completion date", "service date", "performed date",
        "submission date", "submitted date", "submitted at", "date completed",
    ),
    "technician": (
        "technician", "tech", "service technician", "performed by", "completed by", "submitted by",
        "user", "employee",
    ),
    "serial": (
        "serial", "serial number", "serial #", "equipment serial", "equipment serial number",
        "unit serial", "s/n",
    ),
    "manufacturer": ("manufacturer", "make", "equipment make", "equipment manufacturer", "brand"),
    "model": ("model", "model number", "model #", "equipment model", "equipment model number"),
    "complaint": (
        "complaint", "customer complaint", "reported problem", "problem", "problem description",
        "service request", "issue", "reported issue", "description",
    ),
    "finding": (
        "finding", "findings", "diagnosis", "diagnostic", "diagnostic findings", "technician finding",
        "observations", "observed condition",
    ),
    "cause": ("cause", "root cause", "failure cause", "reason"),
    "correction": (
        "correction", "work performed", "work completed", "repair", "repairs", "resolution",
        "completed work", "corrective action", "service performed", "technician notes", "notes",
    ),
    "verification": (
        "verification", "verified", "tested", "test result", "test results", "final test",
        "operational test", "result",
    ),
    "parts_text": ("parts", "parts used", "materials", "materials used", "parts installed"),
    "customer_report": ("customer report", "customer notes", "summary for customer", "service summary"),
}


def _key(value: Any) -> str:
    text = str(value or "").strip().lower().replace("_", " ")
    text = re.sub(r"[\u2013\u2014]", "-", text)
    text = re.sub(r"[^a-z0-9#/()+&. -]+", " ", text)
    return re.sub(r"\s+", " ", text).strip(" :-")


def _text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, (dict, list)):
        rendered = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    else:
        rendered = str(value)
    rendered = rendered.strip()
    return rendered or None


def _flatten(value: Any, prefix: str = "") -> dict[str, Any]:
    out: dict[str, Any] = {}
    if isinstance(value, dict):
        for raw_key, item in value.items():
            k = _key(raw_key)
            path = f"{prefix} {k}".strip() if prefix else k
            if isinstance(item, dict):
                out.update(_flatten(item, path))
            elif isinstance(item, list) and item and all(isinstance(x, dict) for x in item):
                # Preserve structured repeating groups in a compact form while also
                # exposing first-record leaves for common FastField exports.
                out[path] = item
                out.update(_flatten(item[0], path))
            else:
                out[path] = item
                # The leaf key is often what the FastField form designer actually named.
                out.setdefault(k, item)
    return out


def _strip_html(raw: str) -> str:
    raw = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", raw)
    raw = re.sub(r"(?i)<br\s*/?>", "\n", raw)
    raw = re.sub(r"(?i)</(p|div|tr|li|h[1-6])>", "\n", raw)
    raw = re.sub(r"(?s)<[^>]+>", " ", raw)
    raw = html.unescape(raw)
    raw = re.sub(r"[ \t]+", " ", raw)
    return re.sub(r"\n\s*\n+", "\n", raw).strip()


def parse_key_value_text(text: str) -> dict[str, str]:
    """Extract conservative Label: Value pairs from a report/email body.

    Continuation lines append to the prior value. Lines without a clear separator are
    intentionally ignored rather than guessed into fields.
    """
    out: dict[str, str] = {}
    current: str | None = None
    for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        line = line.strip()
        if not line:
            current = None
            continue
        match = re.match(r"^([^:]{2,80}):\s*(.*)$", line)
        if match:
            label = _key(match.group(1))
            value = match.group(2).strip()
            if label and value:
                out[label] = value
                current = label
            else:
                current = None
        elif current and len(line) < 500:
            out[current] = f"{out[current]} {line}".strip()
    return out


def _rows_from_csv(data: bytes) -> list[dict[str, Any]]:
    text = data.decode("utf-8-sig", errors="replace")
    return [dict(row) for row in csv.DictReader(io.StringIO(text))]


def _rows_from_json(data: bytes) -> list[dict[str, Any]]:
    value = json.loads(data.decode("utf-8-sig"))
    if isinstance(value, dict):
        # Common API/export wrappers: {data:[...]}, {submission:{...}}, etc.
        for key in ("data", "submissions", "records", "rows", "results"):
            nested = value.get(key)
            if isinstance(nested, list) and nested and all(isinstance(x, dict) for x in nested):
                return [dict(x) for x in nested]
        for key in ("submission", "record", "result"):
            nested = value.get(key)
            if isinstance(nested, dict):
                return [dict(nested)]
        return [value]
    if isinstance(value, list):
        return [dict(x) for x in value if isinstance(x, dict)]
    return []


def _rows_from_xml(data: bytes) -> list[dict[str, Any]]:
    root = ET.fromstring(data)

    def leaves(node: ET.Element) -> dict[str, str]:
        result: dict[str, str] = {}
        for child in list(node):
            if list(child):
                for key, value in leaves(child).items():
                    result.setdefault(key, value)
            elif child.text and child.text.strip():
                result[_key(child.tag.split("}")[-1])] = child.text.strip()
        return result

    children = list(root)
    if len(children) > 1 and len({_key(c.tag.split("}")[-1]) for c in children}) == 1:
        rows = [leaves(c) for c in children]
        if any(rows):
            return rows
    return [leaves(root)]


def _rows_from_xlsx(data: bytes) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency gate
        raise RuntimeError("Excel intake support is not installed") from exc
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = book[book.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    try:
        headers = [str(x or "").strip() for x in next(rows)]
    except StopIteration:
        return []
    result = []
    for row in rows:
        if not any(v not in (None, "") for v in row):
            continue
        result.append({headers[i] or f"column_{i+1}": value for i, value in enumerate(row) if i < len(headers)})
    return result


def _text_from_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency gate
        raise RuntimeError("PDF intake support is not installed") from exc
    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _text_from_docx(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - dependency gate
        raise RuntimeError("Word intake support is not installed") from exc
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(": ".join(cells[:2]) if len(cells) == 2 else " | ".join(cells))
    return "\n".join(parts)


def attachment_records(attachment: MailAttachment) -> list[dict[str, Any]]:
    name = attachment.name.lower()
    ctype = (attachment.content_type or mimetypes.guess_type(name)[0] or "").lower()
    if name.endswith(".json") or ctype == "application/json":
        return _rows_from_json(attachment.data)
    if name.endswith(".csv") or ctype in ("text/csv", "application/csv"):
        return _rows_from_csv(attachment.data)
    if name.endswith(".xml") or ctype in ("application/xml", "text/xml"):
        return _rows_from_xml(attachment.data)
    if name.endswith(".xlsx") or ctype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        return _rows_from_xlsx(attachment.data)
    if name.endswith(".pdf") or ctype == "application/pdf":
        return [parse_key_value_text(_text_from_pdf(attachment.data))]
    if name.endswith(".docx") or ctype == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return [parse_key_value_text(_text_from_docx(attachment.data))]
    if ctype.startswith("text/") or name.endswith(".txt"):
        return [parse_key_value_text(attachment.data.decode("utf-8", errors="replace"))]
    return []


def _find(flat: dict[str, Any], aliases: Iterable[str]) -> str | None:
    normalized = {_key(k): v for k, v in flat.items()}
    for alias in aliases:
        if alias in normalized:
            value = _text(normalized[alias])
            if value:
                return value
    # Nested exports frequently prefix the form/section name. Match only an exact
    # final leaf token, not arbitrary substring, to avoid surprising guesses.
    for alias in aliases:
        suffix = " " + alias
        candidates = [(k, v) for k, v in normalized.items() if k.endswith(suffix)]
        if len(candidates) == 1:
            value = _text(candidates[0][1])
            if value:
                return value
    return None


def normalize_record(record: dict[str, Any], *, body_values: dict[str, str] | None = None) -> dict[str, Any]:
    flat = _flatten(record)
    if body_values:
        for key, value in body_values.items():
            flat.setdefault(key, value)
    normalized: dict[str, Any] = {field: _find(flat, aliases) for field, aliases in ALIASES.items()}
    # Keep a small evidence map for review without dumping giant encoded objects.
    evidence: dict[str, str] = {}
    for key, value in flat.items():
        text = _text(value)
        if text and len(text) <= 800 and len(evidence) < 80:
            evidence[_key(key)] = text
    normalized["extra_fields"] = evidence
    normalized["source_system"] = "FastField"
    return normalized


def parse_eml(raw: bytes) -> MailEnvelope:
    message = BytesParser(policy=policy.default).parsebytes(raw)
    body = ""
    try:
        part = message.get_body(preferencelist=("plain", "html"))
    except AttributeError:
        part = None
    if part is not None:
        content = part.get_content()
        body = str(content)
        if part.get_content_type() == "text/html":
            body = _strip_html(body)
    elif message.get_payload():
        body = str(message.get_payload())
    attachments: list[MailAttachment] = []
    for part in message.iter_attachments():
        data = part.get_payload(decode=True) or b""
        attachments.append(MailAttachment(part.get_filename() or "attachment.bin", part.get_content_type(), data))
    internet_id = str(message.get("Message-ID") or "").strip() or None
    external = internet_id or "eml-" + hashlib.sha256(raw).hexdigest()
    return MailEnvelope(
        external_id=external,
        internet_message_id=internet_id,
        sender=str(message.get("From") or "").strip() or None,
        subject=str(message.get("Subject") or "").strip(),
        received_at=str(message.get("Date") or "").strip() or None,
        body=body,
        attachments=attachments,
    )


def extract_fastfield_records(envelope: MailEnvelope) -> tuple[list[dict[str, Any]], str, str]:
    """Return normalized records and the evidence source used.

    Prefer machine-readable attachments. PDF/Word/email text are fallback sources.
    All source attachments can still be retained independently by ingest_mail_message.
    """
    body_values = parse_key_value_text(envelope.body)
    priority = (".json", ".csv", ".xml", ".xlsx", ".pdf", ".docx", ".txt")
    by_priority = sorted(
        envelope.attachments,
        key=lambda a: next((i for i, suffix in enumerate(priority) if a.name.lower().endswith(suffix)), len(priority)),
    )
    for attachment in by_priority:
        try:
            rows = attachment_records(attachment)
        except Exception:
            continue
        rows = [r for r in rows if isinstance(r, dict) and r]
        if rows:
            suffix = Path(attachment.name).suffix.lower()
            method = "DETERMINISTIC_STRUCTURED" if suffix in (".json", ".csv", ".xml", ".xlsx") else "DETERMINISTIC_RULE"
            normalized = [normalize_record(r, body_values=body_values) for r in rows]
            for item in normalized:
                item["_provenance"] = {
                    "method": method,
                    "source": attachment.name,
                    "human_review_required": True,
                    "ai_used": False,
                }
            return normalized, attachment.name, method
    if body_values:
        normalized = normalize_record(body_values)
        normalized["_provenance"] = {
            "method": "DETERMINISTIC_RULE",
            "source": "email body",
            "human_review_required": True,
            "ai_used": False,
        }
        return [normalized], "email body", "DETERMINISTIC_RULE"
    return [], "none", "DETERMINISTIC_RULE"


def _store_attachment(conn: Any, organization_id: str, actor_user_id: str | None, message_id: str, attachment: MailAttachment) -> str | None:
    if not attachment.data or len(attachment.data) > 25 * 1024 * 1024:
        return None
    digest = hashlib.sha256(attachment.data).hexdigest()
    existing = conn.execute(
        "SELECT id FROM file_records WHERE organization_id=? AND entity_type='mail_intake' AND entity_id=? AND sha256=?",
        (organization_id, message_id, digest),
    ).fetchone()
    if existing:
        return str(existing["id"])
    fid = new_id("file")
    suffix = Path(attachment.name).suffix[:12]
    org_dir = FILES_DIR / organization_id
    org_dir.mkdir(parents=True, exist_ok=True)
    stored = f"{fid}{suffix}"
    (org_dir / stored).write_bytes(attachment.data)
    conn.execute(
        """INSERT INTO file_records(id,organization_id,entity_type,entity_id,category,original_name,stored_name,mime_type,size_bytes,sha256,created_by_user_id,created_at)
           VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",
        (fid, organization_id, "mail_intake", message_id, "FASTFIELD_SOURCE", attachment.name, stored,
         attachment.content_type, len(attachment.data), digest, actor_user_id, utcnow()),
    )
    return fid


def _match_record(conn: Any, organization_id: str, normalized: dict[str, Any]) -> tuple[str, str | None, str | None]:
    customer_name = _text(normalized.get("customer"))
    serial = _text(normalized.get("serial"))
    issue: str | None = None
    customer = None
    equipment = None
    state = "READY"
    if not customer_name:
        return "ISSUE", "FastField submission did not include a recognizable customer name", None
    customer = conn.execute(
        "SELECT id,name FROM customers WHERE organization_id=? AND lower(trim(name))=lower(trim(?))",
        (organization_id, customer_name),
    ).fetchone()
    if not customer:
        return "REVIEW", f"Customer not found: {customer_name}", None
    if serial:
        equipment = conn.execute(
            "SELECT id FROM equipment WHERE organization_id=? AND customer_id=? AND lower(trim(serial_number))=lower(trim(?))",
            (organization_id, customer["id"], serial),
        ).fetchone()
        if not equipment:
            state, issue = "REVIEW", f"Equipment serial not found for this customer: {serial}"
    normalized["equipment_id"] = equipment["id"] if equipment else None
    normalized["matched_customer_id"] = customer["id"]
    return state, issue, customer["id"]


def ingest_mail_message(
    organization_id: str,
    actor_user_id: str | None,
    provider: str,
    envelope: MailEnvelope,
) -> dict[str, Any]:
    """Deduplicate, preserve source evidence, extract FastField data, and stage review."""
    now = utcnow()
    with connect() as conn:
        existing = conn.execute(
            "SELECT * FROM mail_intake_messages WHERE organization_id=? AND provider=? AND external_id=?",
            (organization_id, provider, envelope.external_id),
        ).fetchone()
        if existing:
            return {"message_id": existing["id"], "duplicate": True, "batch_id": existing["import_batch_id"], "state": existing["state"]}
        message_id = new_id("mail")
        conn.execute(
            """INSERT INTO mail_intake_messages(id,organization_id,provider,external_id,internet_message_id,sender,subject,received_at,body_preview,source_kind,state,created_at,updated_at)
               VALUES(?,?,?,?,?,?,?,?,?,'FASTFIELD','RECEIVED',?,?)""",
            (message_id, organization_id, provider, envelope.external_id, envelope.internet_message_id,
             envelope.sender, envelope.subject, envelope.received_at, envelope.body[:2000], now, now),
        )
        for attachment in envelope.attachments:
            _store_attachment(conn, organization_id, actor_user_id, message_id, attachment)

        records, source_name, derivation_method = extract_fastfield_records(envelope)
        if not records:
            conn.execute(
                "UPDATE mail_intake_messages SET state='NEEDS_REVIEW',error=?,updated_at=? WHERE id=?",
                ("No recognizable structured FastField fields were found. Original email and attachments were preserved.", now, message_id),
            )
            return {"message_id": message_id, "duplicate": False, "batch_id": None, "state": "NEEDS_REVIEW", "records": 0}

        batch_id = new_id("import")
        conn.execute(
            """INSERT INTO import_batches(id,organization_id,source_type,filename,status,row_count,ready_count,issue_count,created_by_user_id,created_at)
               VALUES(?,?,?,?,?,?,?,?,?,?)""",
            (batch_id, organization_id, "FASTFIELD_EMAIL", source_name, "STAGED", len(records), 0, 0, actor_user_id, now),
        )
        ready = issues = 0
        for index, normalized in enumerate(records, start=1):
            normalized["source_system"] = "FastField"
            normalized["mail_intake_message_id"] = message_id
            if not normalized.get("source_reference"):
                normalized["source_reference"] = envelope.internet_message_id or envelope.external_id
            state, issue, customer_id = _match_record(conn, organization_id, normalized)
            if state == "READY":
                ready += 1
            else:
                issues += 1
            conn.execute(
                """INSERT INTO import_rows(id,batch_id,row_number,raw_json,normalized_json,state,issue,matched_customer_id,derivation_method,review_state)
                   VALUES(?,?,?,?,?,?,?,?,?,?)""",
                (new_id("irow"), batch_id, index, json.dumps({"source": source_name, "subject": envelope.subject}),
                 json.dumps(normalized), state, issue, customer_id, derivation_method, "AWAITING_HUMAN_REVIEW"),
            )
        conn.execute("UPDATE import_batches SET ready_count=?,issue_count=? WHERE id=?", (ready, issues, batch_id))
        message_state = "AWAITING_HUMAN_REVIEW" if issues == 0 else "NEEDS_REVIEW"
        conn.execute(
            "UPDATE mail_intake_messages SET import_batch_id=?,state=?,updated_at=? WHERE id=?",
            (batch_id, message_state, now, message_id),
        )
    return {
        "message_id": message_id,
        "duplicate": False,
        "batch_id": batch_id,
        "state": message_state,
        "records": len(records),
        "ready_count": ready,
        "issue_count": issues,
        "source": source_name,
    }

_SERVICE_DAY_CODE_RE = re.compile(r"\[SSDAY:([A-Z0-9]{6,12})\]", re.IGNORECASE)


def process_service_day_reply(
    organization_id: str,
    envelope: MailEnvelope,
) -> dict[str, Any] | None:
    """Apply an explicit customer email reply to a service-day confirmation.

    This parser is deliberately narrow. It only acts when a ServiceSlate reply code is
    present and the customer uses unambiguous confirmation/reschedule language. Anything
    else remains ordinary mail for human review rather than being guessed.
    """
    match = _SERVICE_DAY_CODE_RE.search(f"{envelope.subject}\n{envelope.body}")
    if not match:
        return None
    code = match.group(1).upper()
    raw_body = (envelope.body or "").replace("\r\n", "\n")
    # Inspect the customer's new reply, not quoted copies of the message ServiceSlate sent.
    # Outlook/email clients commonly append the original instructions, which themselves contain
    # words such as "DIFFERENT DAY" and must never be mistaken for the customer's decision.
    reply_body = raw_body
    for pattern in (
        r"(?im)^-{2,}\s*original message\s*-{2,}$",
        r"(?im)^from:\s+",
        r"(?im)^on .{0,160} wrote:\s*$",
        r"(?im)^please reply confirm\b",
        r"(?im)^your service with .+ is scheduled for\b",
    ):
        hit = re.search(pattern, reply_body)
        if hit and hit.start() > 0:
            reply_body = reply_body[:hit.start()]
    reply_body = "\n".join(line for line in reply_body.splitlines() if not line.lstrip().startswith(">"))
    body = re.sub(r"\s+", " ", reply_body).strip().lower()
    # Prefer an explicit change request over confirmation only within the customer's new reply.
    change_terms = (
        "different day", "different date", "change the day", "change date",
        "change the date", "reschedule", "another day", "another date",
    )
    confirm_terms = (
        "confirm", "confirmed", "that day works", "date works", "day works",
        "yes that works", "yes, that works", "yes this works",
    )
    decision = "CHANGE_REQUESTED" if any(term in body for term in change_terms) else (
        "CONFIRMED" if any(term in body for term in confirm_terms) else None
    )
    with connect() as conn:
        prior = conn.execute(
            "SELECT state,error FROM mail_intake_messages WHERE organization_id=? AND provider='SERVICE_DAY_REPLY' AND external_id=?",
            (organization_id, envelope.external_id),
        ).fetchone()
        if prior:
            return {"matched": True, "state": prior["state"], "reply_code": code, "already_processed": True, "reason": prior["error"]}
        row = conn.execute(
            """SELECT sc.*,c.name customer_name,c.email customer_email
               FROM service_day_confirmations sc JOIN customers c ON c.id=sc.customer_id
               WHERE sc.organization_id=? AND upper(sc.reply_code)=? ORDER BY sc.created_at DESC LIMIT 1""",
            (organization_id, code),
        ).fetchone()
        if not row:
            return {"matched": False, "reason": "unknown_reply_code", "reply_code": code}
        now = utcnow()
        def remember(state: str, reason: str | None = None) -> None:
            conn.execute(
                """INSERT INTO mail_intake_messages(id,organization_id,provider,external_id,internet_message_id,sender,subject,received_at,body_preview,source_kind,state,error,created_at,updated_at)
                   VALUES(?,?,?,?,?,?,?,?,?,'SERVICE_DAY_REPLY',?,?,?,?)""",
                (new_id("mail"), organization_id, "SERVICE_DAY_REPLY", envelope.external_id, envelope.internet_message_id, envelope.sender, envelope.subject, envelope.received_at, (envelope.body or "")[:1000], state, reason, now, now),
            )
        if row["status"] == "STALE":
            remember("STALE", "scheduled_date_changed")
            return {"matched": True, "state": "STALE", "confirmation_id": row["id"], "reply_code": code}
        sender = parseaddr(envelope.sender or "")[1].strip().lower()
        expected_raw = row["recipient"] or row["customer_email"] or ""
        expected = (parseaddr(expected_raw)[1] or str(expected_raw)).strip().lower()
        if expected and sender != expected:
            # A different sender may be a legitimate alternate contact; a staff member must review it.
            remember("NEEDS_HUMAN_REVIEW", "sender_does_not_match_recipient")
            return {
                "matched": True,
                "state": "NEEDS_HUMAN_REVIEW",
                "confirmation_id": row["id"],
                "reply_code": code,
                "reason": "sender_does_not_match_recipient",
            }
        if decision is None:
            remember("NEEDS_HUMAN_REVIEW", "ambiguous_reply")
            return {
                "matched": True,
                "state": "NEEDS_HUMAN_REVIEW",
                "confirmation_id": row["id"],
                "reply_code": code,
                "reason": "ambiguous_reply",
            }
        if row["status"] == "CONFIRMED" and decision == "CONFIRMED":
            remember("CONFIRMED", "already_confirmed")
            return {"matched": True, "state": "CONFIRMED", "confirmation_id": row["id"], "reply_code": code, "already_confirmed": True}
        remember(decision)
        if decision == "CONFIRMED":
            conn.execute(
                """UPDATE service_day_confirmations SET status='CONFIRMED',channel='EMAIL_REPLY',
                   confirmed_by_name=?,confirmed_at=?,reply_received_at=?,updated_at=? WHERE id=?""",
                (envelope.sender or row["customer_name"], now, now, now, row["id"]),
            )
        else:
            conn.execute(
                """UPDATE service_day_confirmations SET status='CHANGE_REQUESTED',channel='EMAIL_REPLY',
                   change_request_note=?,reply_received_at=?,updated_at=? WHERE id=?""",
                ((envelope.body or "")[:1000], now, now, row["id"]),
            )
        # Customer reply is human-authored evidence, so it may directly establish the customer's decision.
        conn.execute(
            """INSERT INTO audit_events(organization_id,actor_user_id,entity_type,entity_id,action,summary,created_at)
               VALUES(?,NULL,'job',?,?,?,?)""",
            (
                organization_id,
                row["job_id"],
                f"CUSTOMER_SERVICE_DAY_{decision}",
                f"Customer email reply {decision.lower().replace('_',' ')} service day {row['scheduled_date']} [{code}]",
                now,
            ),
        )
    return {"matched": True, "state": decision, "confirmation_id": row["id"], "reply_code": code}
